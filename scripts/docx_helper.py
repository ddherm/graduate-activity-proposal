#!/usr/bin/env python3
"""Inspect and safely fill explicit fields in DOCX templates; not a proposal generator.

Requires: python-docx. Patch creates a candidate; commit replaces an explicitly named
source after caller-performed visual/semantic review and an internal backup.
"""
from __future__ import annotations
import argparse
import hashlib
from io import BytesIO
import json
import os
from pathlib import Path
import re
import stat
import sys
import tempfile
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

SCHEMA = 1
PROTECTED = re.compile(r"签名|签字|公章|团委[\s\n]*意见|研究生会意见|参与院研会意见")
UNFILLED = [
    ("placeholder", re.compile(r"(?<![A-Za-z])[XＸxｘ]{2,}|20[2XＸxｘ][XＸxｘ]|[XＸxｘ](?=学院|月|日|个|人|篇|条|元|活动|特色)|(?<=[：:\t])[XＸxｘ](?=\s|$)|^[XＸxｘ]$")),
    ("instruction", re.compile(r"方正仿宋|根据内容调整大小|最多可填写|可自行添加|正式表格请删除|若参与学院|不得少于500字|本项目的活动主题和意义|明确写出预期活动成效|全体协办研究生会|填写及签字人员|需加盖同级团组织|未涉及信息|【牵头学院】")),
    ("draft_marker", re.compile(r"待补充|待确认|待填写|TODO|TBD|示例数据|教学示例|模拟数据")),
]
UNSAFE_TAGS = {
    "drawing", "pict", "object", "fldChar", "instrText", "hyperlink", "sdt",
    "footnoteReference", "endnoteReference", "commentReference", "del", "ins",
    "sym", "lastRenderedPageBreak",
}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def bytes_digest(data):
    return hashlib.sha256(data).hexdigest()


def format_xml(element):
    return etree.tostring(element, encoding="unicode") if element is not None else None


def run_format_key(run):
    props = run._r.rPr
    return etree.tostring(props, method="c14n") if props is not None else b""


def load_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def dump(value):
    print(json.dumps(value, ensure_ascii=False, indent=2))


def inventory(doc):
    entries, objects = [], {}

    def paragraph(p, pid, parent=None):
        data = {"id": pid, "type": "paragraph", "text": p.text,
                "style": p.style.name if p.style is not None else None,
                "ppr_xml": format_xml(p._p.pPr), "run_ids": [],
                "protected": bool(PROTECTED.search(p.text))}
        if parent:
            data["parent"] = parent
        entries.append(data)
        objects[pid] = p
        for ri, run in enumerate(p.runs):
            rid = f"{pid}.run{ri}"
            data["run_ids"].append(rid)
            entries.append({"id": rid, "type": "run", "parent": pid,
                            "text": run.text, "rpr_xml": format_xml(run._r.rPr),
                            "protected": bool(PROTECTED.search(run.text))})
            objects[rid] = run
        return pid

    def table(t, tid):
        entries.append({"id": tid, "type": "table", "rows": len(t.rows),
                        "grid_columns": len(t.columns)})
        objects[tid] = t
        seen = {}
        for ri, row in enumerate(t.rows):
            for ci, cell in enumerate(row.cells):
                cid = f"{tid}.r{ri}.c{ci}"
                if cell._tc in seen:
                    seen[cell._tc]["aliases"].append(cid)
                    continue
                data = {"id": cid, "type": "cell", "text": cell.text,
                        "aliases": [], "paragraph_ids": [],
                        "protected": bool(PROTECTED.search(cell.text)),
                        "nested_tables": len(cell.tables)}
                seen[cell._tc] = data
                entries.append(data)
                objects[cid] = cell
                for pi, p in enumerate(cell.paragraphs):
                    data["paragraph_ids"].append(paragraph(p, f"{cid}.p{pi}", cid))
                for ti, nested in enumerate(cell.tables):
                    table(nested, f"{cid}.t{ti}")

    for i, p in enumerate(doc.paragraphs):
        paragraph(p, f"body.p{i}")
    for i, t in enumerate(doc.tables):
        table(t, f"body.t{i}")
    return entries, objects


def inspect_file(path):
    snapshot = Path(path).read_bytes()
    source_hash = bytes_digest(snapshot)
    doc = Document(BytesIO(snapshot))
    entries, _ = inventory(doc)
    return {"schema_version": SCHEMA, "source": str(Path(path).resolve()),
            "source_sha256": source_hash, "coordinate_base": 0,
            "scope": "Main document body, tables and nested tables. Merged cells have one canonical id and alias coordinates. Header/footer/textbox fields require separate review.",
            "format_note": "ppr_xml/rpr_xml show direct properties only; effective formatting can inherit from styles. Run ids support exact text replacement without collapsing sibling runs.",
            "patch_example": {"schema_version": SCHEMA, "source_sha256": source_hash,
                              "operations": [{"target": "body.p1", "expect": "copy exact text from entries", "text": "new text"}]},
            "entries": entries}


def validate_text(value):
    if not isinstance(value, str):
        raise ValueError("text / paragraphs values must be strings")
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise ValueError("replacement contains XML-incompatible control characters")


def check_paragraph(p):
    allowed_paragraph_children = {"pPr", "r", "bookmarkStart", "bookmarkEnd", "proofErr"}
    allowed_run_children = {"rPr", "t", "tab", "br", "cr", "noBreakHyphen", "softHyphen"}
    for child in p._p:
        local = etree.QName(child).localname
        if child.tag not in {qn("w:" + x) for x in allowed_paragraph_children}:
            raise ValueError(f"paragraph contains unsupported {local}; use a document editor")
        if child.tag == qn("w:r"):
            for run_child in child:
                if run_child.tag not in {qn("w:" + x) for x in allowed_run_children}:
                    raise ValueError(f"run contains unsupported {etree.QName(run_child).localname}; use a document editor")
    for node in p._p.iter():
        local = etree.QName(node).localname
        if local in UNSAFE_TAGS:
            raise ValueError(f"paragraph contains {local}; use a document editor for this field")
        if local == "br" and node.get(qn("w:type"), "textWrapping") != "textWrapping":
            raise ValueError("paragraph contains a page/column break; use a document editor")


def put_text(p, text):
    """Replace a uniform-format paragraph; mixed direct run formats are refused."""
    check_paragraph(p)
    check_uniform_format(p)
    runs = list(p.runs)
    first = next((r for r in runs if r.text), runs[0] if runs else None)
    if first is None:
        first = p.add_run()
    for run in runs:
        for child in list(run._r):
            if child.tag in {qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr"), qn("w:noBreakHyphen"), qn("w:softHyphen")}:
                run._r.remove(child)
    first.text = text


def check_uniform_format(p):
    runs = list(p.runs)
    relevant = [r for r in runs if r.text] or runs
    if len({run_format_key(r) for r in relevant}) > 1:
        raise ValueError("paragraph has mixed direct run formatting; use exact .runN targets or a document editor")


def protected_target(target, metadata):
    while target:
        meta = metadata[target]
        if meta["protected"]:
            return True
        target = meta.get("parent")
    return False


def validate_docx_bytes(data):
    """Structural/CRC validation only; never a layout or content approval."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            names = archive.namelist()
            if len(names) != len(set(names)):
                raise ValueError("DOCX contains duplicate ZIP members")
            if not {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= set(names):
                raise ValueError("candidate is not a complete DOCX package")
            bad_member = archive.testzip()
            if bad_member:
                raise ValueError(f"DOCX CRC validation failed for {bad_member}")
            etree.fromstring(archive.read("word/document.xml"))
        Document(BytesIO(data))
    except (BadZipFile, KeyError, etree.XMLSyntaxError, EOFError, RuntimeError) as exc:
        raise ValueError(f"invalid DOCX package: {exc}") from exc


def check_file(path):
    findings = []
    with ZipFile(path) as z:
        for name in z.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(z.read(name))
            for p in root.iter(qn("w:p")):
                text = "".join((n.text or "") if n.tag == qn("w:t") else "\t" if n.tag == qn("w:tab") else "\n"
                               for n in p.iter() if n.tag in {qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr")})
                matches = []
                for kind, pattern in UNFILLED:
                    values = list(dict.fromkeys(m.group(0) for m in pattern.finditer(text)))
                    if values:
                        matches.append({"kind": kind, "matches": values})
                if matches:
                    findings.append({"part": name, "xml_path": root.getroottree().getpath(p),
                                     "text": text, "protected_manual_field": bool(PROTECTED.search(text)), "matches": matches})
    return {"source": str(Path(path).resolve()), "findings_count": len(findings), "findings": findings,
            "meaning": "Heuristic review leads only. Findings can include legitimate instructions or manual signature fields; absence does not establish completeness or submission readiness. Blank fields, facts, totals, page layout and policy compliance need separate review."}


def patch_file(source, specpath, output):
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("patch input and candidate must differ; use commit after rendering and reviewing the candidate")
    if output.exists():
        raise ValueError("output already exists; choose a new path (no overwrite)")
    spec = load_json(specpath)
    if spec.get("schema_version") != SCHEMA:
        raise ValueError(f"schema_version must be {SCHEMA}")
    snapshot = source.read_bytes()
    source_hash = bytes_digest(snapshot)
    if spec.get("source_sha256") != source_hash:
        raise ValueError("source_sha256 mismatch; inspect this exact source again")
    operations = spec.get("operations")
    if not isinstance(operations, list) or not operations:
        raise ValueError("operations must be a non-empty list")
    doc = Document(BytesIO(snapshot))
    entries, objects = inventory(doc)
    metadata = {e["id"]: e for e in entries}
    planned, used_paragraphs, used_runs = [], set(), {}
    for n, op in enumerate(operations):
        if not isinstance(op, dict) or not {"target", "expect"} <= op.keys():
            raise ValueError(f"operation {n}: target and expect are required")
        if set(op) - {"target", "expect", "text", "paragraphs"}:
            raise ValueError(f"operation {n}: unsupported keys")
        target = op["target"]
        if target not in objects or metadata[target]["type"] == "table":
            raise ValueError(f"operation {n}: unknown/non-editable target {target}; use canonical inspect id")
        obj, meta = objects[target], metadata[target]
        if not isinstance(op["expect"], str) or obj.text != op["expect"]:
            raise ValueError(f"operation {n}: expect mismatch at {target}; no file was written")
        if protected_target(target, metadata):
            raise ValueError(f"operation {n}: {target} is a manual opinion/signature/seal area; not modified by this helper")
        if ("text" in op) == ("paragraphs" in op):
            raise ValueError(f"operation {n}: provide exactly one of text or paragraphs")
        if meta["type"] == "run":
            if "paragraphs" in op:
                raise ValueError(f"operation {n}: run target requires text")
            pid = meta["parent"]
            if pid in used_paragraphs or target in used_runs.get(pid, set()):
                raise ValueError(f"operation {n}: overlapping target {target}")
            used_runs.setdefault(pid, set()).add(target)
            validate_text(op["text"])
            check_paragraph(objects[pid])
            planned.append(("run", obj, op["text"]))
            continue
        if meta["type"] == "paragraph":
            if "paragraphs" in op:
                raise ValueError(f"operation {n}: paragraph target requires text")
            changes = [(target, obj, op["text"])]
        else:
            if obj.tables:
                raise ValueError(f"operation {n}: cell contains nested tables; edit its paragraphs individually")
            if "text" in op:
                if len(obj.paragraphs) != 1:
                    raise ValueError(f"operation {n}: multi-paragraph cell requires paragraphs array of exact existing length, or individual paragraph operations")
                texts = [op["text"]]
            else:
                texts = op["paragraphs"]
                if not isinstance(texts, list) or len(texts) != len(obj.paragraphs):
                    raise ValueError(f"operation {n}: paragraphs count must equal {len(obj.paragraphs)}")
            changes = [(pid, p, text) for pid, p, text in zip(meta["paragraph_ids"], obj.paragraphs, texts)]
        for pid, p, new in changes:
            if pid in used_paragraphs or pid in used_runs:
                raise ValueError(f"operation {n}: overlapping target {pid}")
            used_paragraphs.add(pid)
            validate_text(new)
            check_paragraph(p)
            check_uniform_format(p)
            planned.append(("paragraph", p, new))
    for kind, obj, new in planned:
        if kind == "run":
            obj.text = new  # python-docx preserves this run's rPr and sibling runs.
        else:
            put_text(obj, new)
    # Preserve every other package part exactly, including images, styles and metadata.
    xml = etree.tostring(doc._element, xml_declaration=True, encoding="UTF-8", standalone=True)
    if not output.parent.is_dir():
        raise ValueError("output parent directory must already exist")
    candidate = BytesIO()
    with ZipFile(BytesIO(snapshot)) as src, ZipFile(candidate, "w") as out:
        out.comment = src.comment
        for info in src.infolist():
            out.writestr(info, xml if info.filename == "word/document.xml" else src.read(info.filename))
    data = candidate.getvalue()
    validate_docx_bytes(data)
    if digest(source) != source_hash:
        raise ValueError("source changed while preparing candidate; no file was written")
    # Exclusive creation avoids overwriting a concurrently-created file.
    created = False
    try:
        with output.open("xb") as out_handle:
            created = True
            out_handle.write(data)
            out_handle.flush()
            os.fsync(out_handle.fileno())
    except OSError:
        if created:
            output.unlink(missing_ok=True)
        raise
    return {"output": str(output.resolve()), "source_snapshot_sha256": source_hash,
            "candidate_sha256": bytes_digest(data), "operations_applied": len(operations),
            "paragraphs_changed": len(used_paragraphs | set(used_runs)),
            "review_required": "Run check, compare fields and totals, then render and inspect every page before commit. The helper does not refresh TOC/page numbers, remove rows, synthesize signatures, or certify submission readiness."}


def commit_file(source, candidate, source_sha256, candidate_sha256, backup_dir):
    """Commit caller-reviewed bytes atomically at source, with a verified backup.

    This function does not render files or certify that the caller reviewed them.
    Hash/format/write failures before os.replace leave source untouched.
    """
    source, candidate, backup_dir = Path(source), Path(candidate), Path(backup_dir)
    if source.is_symlink():
        raise ValueError("commit refuses a symlink source; use the actual document path")
    if source.resolve() == candidate.resolve():
        raise ValueError("source and candidate must differ")
    for label, expected in (("source", source_sha256), ("candidate", candidate_sha256)):
        if not isinstance(expected, str) or not re.fullmatch(r"[0-9a-f]{64}", expected):
            raise ValueError(f"{label}_sha256 must be a lowercase SHA-256 digest")
    with source.open("rb") as handle:
        source_stat = os.fstat(handle.fileno())
        original = handle.read()
    if not stat.S_ISREG(source_stat.st_mode):
        raise ValueError("source must be a regular file")
    reviewed = candidate.read_bytes()
    if bytes_digest(original) != source_sha256:
        raise ValueError("source_sha256 mismatch; original path was not modified")
    if bytes_digest(reviewed) != candidate_sha256:
        raise ValueError("candidate_sha256 mismatch; candidate changed since review; original path was not modified")
    validate_docx_bytes(original)
    validate_docx_bytes(reviewed)

    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_fd, backup_name = tempfile.mkstemp(prefix=f"{source.stem}.{source_sha256[:12]}.",
                                             suffix=".backup.docx", dir=backup_dir)
    backup = Path(backup_name)
    try:
        with os.fdopen(backup_fd, "wb") as handle:
            handle.write(original)
            handle.flush()
            os.fsync(handle.fileno())
        if digest(backup) != source_sha256:
            raise ValueError("backup hash mismatch; original path was not modified")
    except Exception:
        backup.unlink(missing_ok=True)
        raise

    # Same-directory staging is essential when source is on a removable volume.
    temp_fd, temp_name = tempfile.mkstemp(prefix=f".{source.name}.", suffix=".tmp", dir=source.parent)
    staged = Path(temp_name)
    replaced = False
    try:
        with os.fdopen(temp_fd, "wb") as handle:
            handle.write(reviewed)
            os.fchmod(handle.fileno(), stat.S_IMODE(source_stat.st_mode))
            handle.flush()
            os.fsync(handle.fileno())
        if digest(staged) != candidate_sha256:
            raise ValueError("staged candidate hash mismatch; original path was not modified")
        if source.is_symlink() or digest(source) != source_sha256:
            raise ValueError("source changed before commit; original path was not modified")
        if digest(candidate) != candidate_sha256:
            raise ValueError("candidate changed before commit; original path was not modified")
        os.replace(staged, source)
        replaced = True
        if digest(source) != candidate_sha256:
            raise ValueError("post-commit hash mismatch")
    except Exception as exc:
        if replaced:
            raise ValueError(f"replacement occurred but post-commit verification failed; original backup retained at {backup}: {exc}") from exc
        raise
    finally:
        staged.unlink(missing_ok=True)
    return {"output": str(source.resolve()), "committed": True,
            "source_before_sha256": source_sha256, "source_after_sha256": candidate_sha256,
            "backup": str(backup.resolve()), "backup_sha256": source_sha256,
            "permissions_preserved": oct(stat.S_IMODE(source_stat.st_mode)),
            "layout_review_verified_by_tool": False,
            "review_note": "Commit validates hashes and DOCX structure only. The caller must render and inspect the exact candidate before invoking commit."}


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("inspect", help="List main-body paragraphs, tables and deduplicated merged cells as JSON")
    p.add_argument("source", type=Path)
    p = sub.add_parser("patch", help="Apply explicit JSON operations to a new DOCX; source hash and exact expect required")
    p.add_argument("source", type=Path)
    p.add_argument("patch_json", type=Path)
    p.add_argument("output", type=Path)
    p = sub.add_parser("commit", help="Replace source with a previously rendered/reviewed candidate; verifies hashes and creates an internal backup")
    p.add_argument("source", type=Path)
    p.add_argument("candidate", type=Path)
    p.add_argument("--source-sha256", required=True)
    p.add_argument("--candidate-sha256", required=True)
    p.add_argument("--backup-dir", type=Path, required=True)
    p = sub.add_parser("check", help="Find X/XX/202X, teaching instructions and draft markers in Word XML parts; not a completeness verdict")
    p.add_argument("source", type=Path)
    p.add_argument("--fail-on-findings", action="store_true", help="Exit 1 when review leads exist (default exit 0)")
    args = parser.parse_args()
    try:
        if args.source.name.startswith("._"):
            raise ValueError("AppleDouble ._ metadata files are not DOCX templates")
        if args.command == "inspect":
            result = inspect_file(args.source)
        elif args.command == "patch":
            result = patch_file(args.source, args.patch_json, args.output)
        elif args.command == "commit":
            result = commit_file(args.source, args.candidate, args.source_sha256,
                                 args.candidate_sha256, args.backup_dir)
        else:
            result = check_file(args.source)
        dump(result)
        if args.command == "check" and args.fail_on_findings and result["findings_count"]:
            return 1
        return 0
    except (ValueError, KeyError, OSError, TypeError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
